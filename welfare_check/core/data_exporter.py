"""
데이터 내보내기 모듈 - 엑셀 파일로 내보내기
"""

from typing import List, Dict, Any, Optional
from datetime import datetime

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


# 대상자 내보내기 열 정의
BENEFICIARY_COLUMNS = [
    ("이름", "name"),
    ("전화번호", "phone"),
    ("주소", "address"),
    ("생년월일", "birth_date"),
    ("성별", "gender"),
    ("담당자", "manager"),
    ("서비스", "service_type"),
    ("돌봄등급", "grade"),
    ("비고", "note"),
    ("상태", "status"),
    ("등록일", "registered_date"),
]

# 안부 확인 이력 내보내기 열 정의
CHECK_LOG_COLUMNS = [
    ("날짜", "check_date"),
    ("대상자", "beneficiary_name"),
    ("확인방법", "check_type"),
    ("결과", "result"),
    ("확인자", "checked_by"),
    ("메모", "memo"),
]

# 연결망 내보내기 열 정의
NETWORK_COLUMNS = [
    ("이름", "contact_name"),
    ("관계", "relationship"),
    ("연락처", "phone"),
    ("역할", "role"),
    ("방문주기", "visit_cycle"),
    ("비고", "note"),
]


def _apply_header_style(ws, col_count: int) -> None:
    """헤더 행에 서식 적용"""
    header_font = Font(name="맑은 고딕", bold=True, size=11, color="FFFFFF")
    header_fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    for col in range(1, col_count + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border


def _auto_adjust_width(ws) -> None:
    """열 너비 자동 조절"""
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                val = str(cell.value) if cell.value else ""
                # 한글은 2바이트로 계산
                length = sum(2 if ord(c) > 127 else 1 for c in val)
                max_length = max(max_length, length)
            except Exception:
                pass
        adjusted_width = min(max_length + 4, 50)
        ws.column_dimensions[col_letter].width = adjusted_width


def _write_data_rows(ws, data: List[Dict[str, Any]], columns: List[tuple], start_row: int = 2) -> None:
    """데이터 행 쓰기"""
    data_font = Font(name="맑은 고딕", size=10)
    data_align = Alignment(vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )
    even_fill = PatternFill(start_color="F0F4FA", end_color="F0F4FA", fill_type="solid")

    for row_idx, item in enumerate(data):
        row_num = start_row + row_idx
        for col_idx, (_, field) in enumerate(columns):
            cell = ws.cell(row=row_num, column=col_idx + 1)
            cell.value = item.get(field, "")
            cell.font = data_font
            cell.alignment = data_align
            cell.border = thin_border
            if row_idx % 2 == 1:
                cell.fill = even_fill


def generate_filename(prefix: str = "복지대상자_내보내기") -> str:
    """내보내기 파일명 자동 생성"""
    date_str = datetime.now().strftime("%Y%m%d")
    return f"{prefix}_{date_str}.xlsx"


def export_current_view(
    data: List[Dict[str, Any]],
    file_path: str,
    columns: Optional[List[tuple]] = None
) -> None:
    """
    현재 화면(필터 적용) 데이터를 엑셀로 내보내기

    Args:
        data: 내보낼 데이터 목록
        file_path: 저장할 파일 경로
        columns: 열 정의 (None이면 기본 대상자 열)
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl 라이브러리가 설치되지 않았습니다.")

    if columns is None:
        columns = BENEFICIARY_COLUMNS

    wb = Workbook()
    ws = wb.active
    ws.title = "대상자 목록"

    # 헤더 쓰기
    for col_idx, (header, _) in enumerate(columns):
        ws.cell(row=1, column=col_idx + 1, value=header)

    _apply_header_style(ws, len(columns))
    _write_data_rows(ws, data, columns)
    _auto_adjust_width(ws)

    wb.save(file_path)


def export_by_manager(
    data: List[Dict[str, Any]],
    file_path: str
) -> None:
    """
    담당자별 시트 분리 내보내기

    Args:
        data: 전체 대상자 데이터
        file_path: 저장할 파일 경로
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl 라이브러리가 설치되지 않았습니다.")

    wb = Workbook()
    # 기본 시트 제거
    wb.remove(wb.active)

    # 담당자별 그룹화
    managers: Dict[str, List[Dict[str, Any]]] = {}
    for item in data:
        mgr = item.get("manager", "미지정") or "미지정"
        if mgr not in managers:
            managers[mgr] = []
        managers[mgr].append(item)

    columns = BENEFICIARY_COLUMNS

    for mgr, items in managers.items():
        # 시트 이름에 사용 불가 문자 제거
        sheet_name = mgr[:31].replace("/", "_").replace("\\", "_")
        ws = wb.create_sheet(title=sheet_name)

        # 헤더 쓰기
        for col_idx, (header, _) in enumerate(columns):
            ws.cell(row=1, column=col_idx + 1, value=header)

        _apply_header_style(ws, len(columns))
        _write_data_rows(ws, items, columns)
        _auto_adjust_width(ws)

    if not managers:
        wb.create_sheet(title="데이터 없음")

    wb.save(file_path)


def export_with_networks(
    beneficiaries: List[Dict[str, Any]],
    networks_by_id: Dict[int, List[Dict[str, Any]]],
    file_path: str
) -> None:
    """
    대상자 + 연결망 정보 포함 내보내기

    Args:
        beneficiaries: 대상자 데이터
        networks_by_id: {대상자ID: 연결망 목록} 딕셔너리
        file_path: 저장할 파일 경로
    """
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl 라이브러리가 설치되지 않았습니다.")

    wb = Workbook()

    # 시트1: 대상자 목록
    ws1 = wb.active
    ws1.title = "대상자 목록"
    columns = BENEFICIARY_COLUMNS
    for col_idx, (header, _) in enumerate(columns):
        ws1.cell(row=1, column=col_idx + 1, value=header)
    _apply_header_style(ws1, len(columns))
    _write_data_rows(ws1, beneficiaries, columns)
    _auto_adjust_width(ws1)

    # 시트2: 연결망 목록
    ws2 = wb.create_sheet(title="1인망 연결망")
    net_columns = [("대상자", "beneficiary_name")] + NETWORK_COLUMNS
    for col_idx, (header, _) in enumerate(net_columns):
        ws2.cell(row=1, column=col_idx + 1, value=header)
    _apply_header_style(ws2, len(net_columns))

    all_networks: List[Dict[str, Any]] = []
    for b in beneficiaries:
        b_id = b.get("id")
        nets = networks_by_id.get(b_id, [])
        for n in nets:
            n_copy = dict(n)
            n_copy["beneficiary_name"] = b.get("name", "")
            all_networks.append(n_copy)

    _write_data_rows(ws2, all_networks, net_columns)
    _auto_adjust_width(ws2)

    wb.save(file_path)


def export_check_logs(
    data: List[Dict[str, Any]],
    file_path: str
) -> None:
    """안부 확인 이력 엑셀 내보내기"""
    if not HAS_OPENPYXL:
        raise ImportError("openpyxl 라이브러리가 설치되지 않았습니다.")

    wb = Workbook()
    ws = wb.active
    ws.title = "안부 확인 이력"

    columns = CHECK_LOG_COLUMNS
    for col_idx, (header, _) in enumerate(columns):
        ws.cell(row=1, column=col_idx + 1, value=header)

    _apply_header_style(ws, len(columns))
    _write_data_rows(ws, data, columns)
    _auto_adjust_width(ws)

    wb.save(file_path)
