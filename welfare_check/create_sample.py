#!/usr/bin/env python3
"""
테스트용 샘플 데이터 생성 스크립트
실행하면 welfare_data.db에 샘플 데이터를 INSERT하고,
샘플_복지대상자.xlsx 엑셀 파일도 함께 생성합니다.
"""

import os
import sys
import random
from datetime import datetime, timedelta

# 경로 설정
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from core.database import Database

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def create_sample_data():
    """샘플 데이터 생성"""
    db_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "welfare_data.db")
    db = Database(db_path)

    print("샘플 데이터 생성을 시작합니다...")

    # ==================== 대상자 데이터 ====================
    managers = ["김복지", "박사회", "이돌봄", "최케어"]

    beneficiaries = [
        # 김복지 담당 (6명)
        {"name": "이순자", "phone": "010-1234-5678", "address": "서울시 강남구 역삼동 123",
         "birth_date": "1945-03-15", "gender": "여", "manager": "김복지",
         "service_type": "AI 안부전화", "grade": "중점돌봄", "note": "독거노인, 당뇨"},
        {"name": "박영수", "phone": "010-2345-6789", "address": "서울시 강남구 삼성동 456",
         "birth_date": "1940-07-22", "gender": "남", "manager": "김복지",
         "service_type": "카카오톡 안부확인", "grade": "일반돌봄", "note": ""},
        {"name": "최미영", "phone": "010-3456-7890", "address": "서울시 서초구 반포동 789",
         "birth_date": "1948-11-03", "gender": "여", "manager": "김복지",
         "service_type": "AI 안부전화", "grade": "중점돌봄", "note": "거동불편, 고혈압"},
        {"name": "정한수", "phone": "010-4567-8901", "address": "서울시 강남구 대치동 101",
         "birth_date": "1938-05-20", "gender": "남", "manager": "김복지",
         "service_type": "방문 돌봄", "grade": "긴급돌봄", "note": "치매의심"},
        {"name": "한영희", "phone": "010-5678-9012", "address": "서울시 서초구 서초동 202",
         "birth_date": "1950-09-12", "gender": "여", "manager": "김복지",
         "service_type": "AI 안부전화", "grade": "예방돌봄", "note": ""},
        {"name": "윤기철", "phone": "010-6789-0123", "address": "서울시 강남구 논현동 303",
         "birth_date": "1942-02-28", "gender": "남", "manager": "김복지",
         "service_type": "카카오톡 안부확인", "grade": "일반돌봄", "note": "당뇨관리"},

        # 박사회 담당 (5명)
        {"name": "김옥순", "phone": "010-7890-1234", "address": "서울시 송파구 잠실동 111",
         "birth_date": "1944-06-18", "gender": "여", "manager": "박사회",
         "service_type": "AI 안부전화", "grade": "중점돌봄", "note": "독거노인"},
        {"name": "이만복", "phone": "010-8901-2345", "address": "서울시 송파구 문정동 222",
         "birth_date": "1936-12-05", "gender": "남", "manager": "박사회",
         "service_type": "방문 돌봄", "grade": "긴급돌봄", "note": "장애 3급, 거동불편"},
        {"name": "조순덕", "phone": "010-9012-3456", "address": "서울시 송파구 가락동 333",
         "birth_date": "1947-04-30", "gender": "여", "manager": "박사회",
         "service_type": "AI 안부전화", "grade": "일반돌봄", "note": ""},
        {"name": "강대호", "phone": "010-0123-4567", "address": "서울시 강동구 천호동 444",
         "birth_date": "1941-08-14", "gender": "남", "manager": "박사회",
         "service_type": "카카오톡 안부확인", "grade": "예방돌봄", "note": "고혈압"},
        {"name": "배금자", "phone": "010-1111-2222", "address": "서울시 강동구 명일동 555",
         "birth_date": "1949-01-25", "gender": "여", "manager": "박사회",
         "service_type": "SMS 안부확인", "grade": "일반돌봄", "note": ""},

        # 이돌봄 담당 (4명)
        {"name": "신영자", "phone": "010-2222-3333", "address": "서울시 관악구 신림동 666",
         "birth_date": "1943-10-08", "gender": "여", "manager": "이돌봄",
         "service_type": "AI 안부전화", "grade": "중점돌봄", "note": "독거노인, 우울증"},
        {"name": "오태식", "phone": "010-3333-4444", "address": "서울시 관악구 봉천동 777",
         "birth_date": "1939-03-17", "gender": "남", "manager": "이돌봄",
         "service_type": "방문 돌봄", "grade": "긴급돌봄", "note": "거동불편, 장애 2급"},
        {"name": "임순애", "phone": "010-4444-5555", "address": "서울시 동작구 상도동 888",
         "birth_date": "1946-07-09", "gender": "여", "manager": "이돌봄",
         "service_type": "AI 안부전화", "grade": "일반돌봄", "note": ""},
        {"name": "장보람", "phone": "010-5555-6666", "address": "서울시 동작구 노량진동 999",
         "birth_date": "1952-11-20", "gender": "여", "manager": "이돌봄",
         "service_type": "카카오톡 안부확인", "grade": "예방돌봄", "note": ""},

        # 최케어 담당 (5명)
        {"name": "류정숙", "phone": "010-6666-7777", "address": "서울시 마포구 합정동 100",
         "birth_date": "1945-05-05", "gender": "여", "manager": "최케어",
         "service_type": "AI 안부전화", "grade": "중점돌봄", "note": "독거노인"},
        {"name": "문재호", "phone": "010-7777-8888", "address": "서울시 마포구 상수동 200",
         "birth_date": "1937-09-23", "gender": "남", "manager": "최케어",
         "service_type": "방문 돌봄", "grade": "긴급돌봄", "note": "치매, 거동불편"},
        {"name": "송미란", "phone": "010-8888-9999", "address": "서울시 용산구 이태원동 300",
         "birth_date": "1951-02-14", "gender": "여", "manager": "최케어",
         "service_type": "AI 안부전화", "grade": "일반돌봄", "note": "고혈압"},
        {"name": "안복동", "phone": "010-9999-0000", "address": "서울시 용산구 한남동 400",
         "birth_date": "1940-12-31", "gender": "남", "manager": "최케어",
         "service_type": "카카오톡 안부확인", "grade": "일반돌봄", "note": ""},
        {"name": "홍순이", "phone": "010-1010-2020", "address": "서울시 마포구 서교동 500",
         "birth_date": "1948-06-15", "gender": "여", "manager": "최케어",
         "service_type": "SMS 안부확인", "grade": "예방돌봄", "note": "당뇨관리"},
    ]

    # 대상자 삽입
    for b in beneficiaries:
        b["status"] = "활성"
        db.add_beneficiary(b)

    print(f"  대상자 {len(beneficiaries)}명 생성 완료")

    # ==================== 연결망 데이터 ====================
    all_b = db.get_all_beneficiaries()
    network_data = []

    relationships_pool = [
        ("김영희", "딸", "1차 연락", "매일"),
        ("박철수", "이웃", "2차 연락", "주 2회"),
        ("이통장", "통·반장", "긴급 연락", "주 1회"),
        ("정자원", "자원봉사자", "정기 방문", "주 2회"),
        ("한복지", "복지사", "수시 확인", "주 1회"),
        ("최민수", "아들", "1차 연락", "주 3회"),
        ("강미숙", "이웃", "2차 연락", "주 1회"),
        ("오반장", "통·반장", "긴급 연락", "월 2회"),
        ("김손녀", "손자녀", "2차 연락", "격주"),
        ("이봉사", "자원봉사자", "정기 방문", "주 1회"),
        ("박배우", "배우자", "1차 연락", "매일"),
        ("정형제", "형제", "2차 연락", "월 1회"),
    ]

    notes_pool = [
        "평일 오전만 가능", "주말에만 연락 가능", "언제든지 가능",
        "오후 2시 이후 가능", "출퇴근 시간 제외", "",
    ]

    phone_counter = 1000
    network_count = 0
    for b in all_b:
        # 1~3명의 연결망 할당
        num_contacts = random.randint(1, 3)
        # 일부 대상자는 연결망 0명 (경고 테스트용)
        if b["name"] in ("윤기철", "안복동"):
            continue

        contacts = random.sample(relationships_pool, min(num_contacts, len(relationships_pool)))
        for contact in contacts:
            phone_counter += 1
            net = {
                "beneficiary_id": b["id"],
                "contact_name": contact[0],
                "relationship": contact[1],
                "phone": f"010-{phone_counter:04d}-{random.randint(1000,9999):04d}",
                "role": contact[2],
                "visit_cycle": contact[3],
                "note": random.choice(notes_pool),
            }
            db.add_network(net)
            network_count += 1

    print(f"  연결망 {network_count}건 생성 완료")

    # ==================== 안부 확인 이력 ====================
    check_types = ["전화", "카카오톡", "방문", "SMS"]
    results = ["확인완료", "확인완료", "확인완료", "미응답", "부재중", "음성사서함"]
    memos = [
        "건강 양호, 식사 정상", "컨디션 좋음", "약간 기침", "병원 방문 예정",
        "식사 거름, 주의 필요", "건강 양호", "외출 중", "수면 중이었음",
        "TV 시청 중", "산책 다녀옴", "", "",
    ]

    today = datetime.now()
    log_count = 0
    for day_offset in range(7):
        date = (today - timedelta(days=day_offset)).strftime("%Y-%m-%d")
        # 하루에 4~5건
        selected = random.sample(all_b, min(random.randint(4, 5), len(all_b)))
        for b in selected:
            log = {
                "beneficiary_id": b["id"],
                "check_date": date,
                "check_type": random.choice(check_types),
                "result": random.choice(results),
                "checked_by": b.get("manager", ""),
                "memo": random.choice(memos),
            }
            db.add_check_log(log)
            log_count += 1

    print(f"  안부 확인 이력 {log_count}건 생성 완료")

    # ==================== 엑셀 파일 생성 ====================
    if HAS_OPENPYXL:
        wb = Workbook()
        ws = wb.active
        ws.title = "복지대상자"

        headers = ["성명", "전화번호", "주소", "생년월일", "성별",
                   "담당자", "서비스", "돌봄등급", "비고"]
        header_font = Font(name="맑은 고딕", bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")

        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row_idx, b in enumerate(beneficiaries, 2):
            ws.cell(row=row_idx, column=1, value=b["name"])
            ws.cell(row=row_idx, column=2, value=b["phone"])
            ws.cell(row=row_idx, column=3, value=b["address"])
            ws.cell(row=row_idx, column=4, value=b["birth_date"])
            ws.cell(row=row_idx, column=5, value=b["gender"])
            ws.cell(row=row_idx, column=6, value=b["manager"])
            ws.cell(row=row_idx, column=7, value=b["service_type"])
            ws.cell(row=row_idx, column=8, value=b["grade"])
            ws.cell(row=row_idx, column=9, value=b.get("note", ""))

        # 열 너비 조절
        widths = [10, 15, 30, 12, 6, 10, 15, 10, 20]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[chr(64 + i)].width = w

        excel_path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                  "샘플_복지대상자.xlsx")
        wb.save(excel_path)
        print(f"  샘플 엑셀 파일 생성 완료: {excel_path}")
    else:
        print("  [경고] openpyxl이 없어 엑셀 파일을 생성하지 못했습니다.")

    db.close()
    print("\n샘플 데이터 생성이 완료되었습니다!")
    print(f"  - 대상자: {len(beneficiaries)}명")
    print(f"  - 담당자: {len(managers)}명")
    print(f"  - 연결망: {network_count}건")
    print(f"  - 안부 확인 이력: {log_count}건")


if __name__ == "__main__":
    create_sample_data()
